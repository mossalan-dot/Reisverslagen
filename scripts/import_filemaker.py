#!/usr/bin/env python3
"""Bouwt de gegevens in data/ opnieuw op vanuit de FileMaker-database van Alan
Moss (export als Excel met veldnamen).

Verwacht twee Excel-bestanden:
  1. de reizigers-export met alle gekoppelde tabellen (Reizen, Afbeeldingen,
     Correspondentie, Poëzie, Reisgezellen, Educatie) als portalvelden;
  2. de losse addenda-export (het Id-veld daarin is niet uniek, dus alle regels
     worden behouden).

Gebruik:
    python3 scripts/import_filemaker.py <reizigers.xlsx> <addenda.xlsx>

De database-export zelf hoort niet in de repository; dit script is de
herhaalbare brug van database naar de YAML-gegevens.
"""

import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
import yaml

WORTEL = Path(__file__).resolve().parent.parent
DATA = WORTEL / "data"
BRON = "gemaakt-op-reis-database"

REISTYPE = {
    "Educatie": "educatiereis", "Diplomatie": "diplomatieke reis",
    "Cultuur/vermaak": "plezierreis", "Zaken": "handelsreis",
    "Militair": "militaire reis", "Religie": "pelgrimsreis", "Anders": "overig",
}
REISTYPE_TITEL = {
    "educatiereis": "Educatiereis", "diplomatieke reis": "Diplomatieke reis",
    "plezierreis": "Plezierreis", "handelsreis": "Zakenreis",
    "militaire reis": "Militaire reis", "pelgrimsreis": "Pelgrimsreis",
    "overig": "Reis", "onbekend": "Reis",
}

# Bibliotheeknaam (genormaliseerd) -> (instelling-id, nette naam, plaats, land, type, website)
INSTELLINGEN = {
    "koninklijke bibliotheek": ("kb-nationale-bibliotheek", "KB, nationale bibliotheek", "Den Haag", "Nederland", "bibliotheek", "https://www.kb.nl"),
    "nationale bibliotheek": ("kbr", "KBR (Koninklijke Bibliotheek van België)", "Brussel", "België", "bibliotheek", "https://www.kbr.be"),
    "nationaal archief": ("nationaal-archief", "Nationaal Archief", "Den Haag", "Nederland", "archief", "https://www.nationaalarchief.nl"),
    "universiteitsbibliotheek amsterdam": ("allard-pierson", "Allard Pierson (Universiteit van Amsterdam)", "Amsterdam", "Nederland", "bibliotheek", "https://allardpierson.nl"),
    "tresoar": ("tresoar", "Tresoar", "Leeuwarden", "Nederland", "archief", "https://www.tresoar.nl"),
    "universiteitsbibliotheek leiden": ("ub-leiden", "Universitaire Bibliotheken Leiden", "Leiden", "Nederland", "bibliotheek", "https://www.bibliotheek.universiteitleiden.nl"),
    "hoge raad van adel": ("hoge-raad-van-adel", "Hoge Raad van Adel", "Den Haag", "Nederland", "overig", "https://www.hogeraadvanadel.nl"),
    "zeeuws archief": ("zeeuws-archief", "Zeeuws Archief", "Middelburg", "Nederland", "archief", "https://www.zeeuwsarchief.nl"),
    "museum plantin-moretus": ("museum-plantin-moretus", "Museum Plantin-Moretus", "Antwerpen", "België", "museum", "https://museumplantinmoretus.be"),
    "scheepvaartmuseum": ("het-scheepvaartmuseum", "Het Scheepvaartmuseum", "Amsterdam", "Nederland", "museum", "https://www.hetscheepvaartmuseum.nl"),
    "nederlands scheepvaartmuseum": ("het-scheepvaartmuseum", "Het Scheepvaartmuseum", "Amsterdam", "Nederland", "museum", "https://www.hetscheepvaartmuseum.nl"),
    "collectie six": ("collectie-six", "Collectie Six", "Amsterdam", "Nederland", "particulier", None),
    "regionaal archief alkmaar": ("regionaal-archief-alkmaar", "Regionaal Archief Alkmaar", "Alkmaar", "Nederland", "archief", "https://www.regionaalarchiefalkmaar.nl"),
    "utrechts archief": ("het-utrechts-archief", "Het Utrechts Archief", "Utrecht", "Nederland", "archief", "https://hetutrechtsarchief.nl"),
    "het utrechts archief": ("het-utrechts-archief", "Het Utrechts Archief", "Utrecht", "Nederland", "archief", "https://hetutrechtsarchief.nl"),
    "universiteitsbibliotheek utrecht": ("ub-utrecht", "Universiteitsbibliotheek Utrecht", "Utrecht", "Nederland", "bibliotheek", "https://www.uu.nl/universiteitsbibliotheek"),
    "universiteitsbibliotheek groningen": ("ub-groningen", "Universiteitsbibliotheek Groningen", "Groningen", "Nederland", "bibliotheek", "https://www.rug.nl/library"),
    "erfgoed leiden en omstreken": ("erfgoed-leiden", "Erfgoed Leiden en Omstreken", "Leiden", "Nederland", "archief", "https://www.erfgoedleiden.nl"),
    "stadsarchief amsterdam": ("stadsarchief-amsterdam", "Stadsarchief Amsterdam", "Amsterdam", "Nederland", "archief", "https://www.amsterdam.nl/stadsarchief"),
    "het gelders archief": ("gelders-archief", "Gelders Archief", "Arnhem", "Nederland", "archief", "https://www.geldersarchief.nl"),
    "gelders archief": ("gelders-archief", "Gelders Archief", "Arnhem", "Nederland", "archief", "https://www.geldersarchief.nl"),
    "noord-hollands archief": ("noord-hollands-archief", "Noord-Hollands Archief", "Haarlem", "Nederland", "archief", "https://noord-hollandsarchief.nl"),
    "westfries archief": ("westfries-archief", "Westfries Archief", "Hoorn", "Nederland", "archief", "https://www.westfriesarchief.nl"),
    "groninger archief": ("groninger-archieven", "Groninger Archieven", "Groningen", "Nederland", "archief", "https://www.groningerarchieven.nl"),
    "centraal bureau voor genealogie": ("cbg", "CBG | Centrum voor familiegeschiedenis", "Den Haag", "Nederland", "overig", "https://cbg.nl"),
    "streekarchivariaat noordwest-veluwe": ("streekarchivariaat-noordwest-veluwe", "Streekarchivariaat Noordwest-Veluwe", "Harderwijk", "Nederland", "archief", None),
    "stadsarchief deventer": ("stadsarchief-deventer", "Stadsarchief Deventer", "Deventer", "Nederland", "archief", None),
    "regionaal archief centrum limburg": ("rhcl", "Regionaal Historisch Centrum Limburg", "Maastricht", "Nederland", "archief", "https://www.rhcl.nl"),
    "historisch centrum overijssel": ("collectie-overijssel", "Collectie Overijssel", "Zwolle", "Nederland", "archief", "https://collectieoverijssel.nl"),
    "bibliotheca di archeologia e storia dell’arte": ("bibliotheca-archeologia-storia-arte", "Bibliotheca di Archeologia e Storia dell’Arte", "Rome", "Italië", "bibliotheek", None),
    "kongelige bibliotëk": ("kongelige-bibliotek", "Det Kongelige Bibliotek", "Kopenhagen", "Denemarken", "bibliotheek", "https://www.kb.dk"),
}

instellingen_gebruikt = {}
onbekende_instellingen = set()


def slug(tekst):
    tekst = unicodedata.normalize("NFKD", str(tekst)).encode("ascii", "ignore").decode()
    tekst = re.sub(r"[^a-z0-9]+", "-", tekst.lower()).strip("-")
    return tekst or "onbekend"


def tekst(v):
    if v is None:
        return ""
    return str(v).replace("\x0b", "\n").strip()


def lijst(v):
    """Meerwaardig veld (regels gescheiden door newline) -> lijst."""
    return [s.strip() for s in tekst(v).split("\n") if s.strip()]


def jaar_uit(s):
    m = re.search(r"\b(1[4-9]\d\d|20\d\d)\b", tekst(s))
    return int(m.group(1)) if m else None


def datumveld(s):
    """'03-11-1679' of '1660' -> dict met jaar en (indien zinvol) datum."""
    s = tekst(s)
    if not s:
        return None
    j = jaar_uit(s)
    d = {}
    if j:
        d["jaar"] = j
    m = re.match(r"^(\d{1,2})[-/](\d{1,2})[-/](\d{4})$", s)
    if m and not (m.group(1) == "1" and m.group(2) == "1"):  # 1-1-JJJJ = alleen jaar
        d["datum"] = s
    return d or None


def plaatsdict(plaats, provincie):
    d = {}
    if tekst(plaats):
        d["plaats"] = tekst(plaats)
    if tekst(provincie):
        d["provincie"] = tekst(provincie)
    return d


def _norm_sleutel(s):
    s = unicodedata.normalize("NFKD", str(s).lower()).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip()


_INSTELLINGEN_NORM = {_norm_sleutel(k): v for k, v in INSTELLINGEN.items()}


def instelling_id(bibliotheek, locatie):
    naam = tekst(bibliotheek)
    if not naam:
        return None
    # meerdere instellingen gescheiden door '/': neem de eerste
    hoofd = naam.split("/")[0].strip()
    loc = tekst(locatie).split("/")[0].strip()
    sleutel = _norm_sleutel(hoofd)
    if sleutel in _INSTELLINGEN_NORM:
        iid, nm, pl, land, typ, web = _INSTELLINGEN_NORM[sleutel]
        instellingen_gebruikt[iid] = dict(id=iid, naam=nm, plaats=pl, land=land, type=typ, website=web)
    else:
        iid = slug(hoofd)
        instellingen_gebruikt.setdefault(iid, dict(
            id=iid, naam=hoofd, plaats=loc or None, controle_nodig=True,
            opmerkingen="Automatisch aangemaakt bij import uit de database; gegevens controleren."))
        onbekende_instellingen.add(hoofd)
    return iid


def schrijf(map_, record):
    pad = DATA / map_ / f"{record['id']}.yaml"
    pad.parent.mkdir(parents=True, exist_ok=True)

    def schoon(x):
        if isinstance(x, dict):
            return {k: schoon(v) for k, v in x.items() if schoon(v) not in (None, "", [], {})}
        if isinstance(x, list):
            return [schoon(v) for v in x if schoon(v) not in (None, "", [], {})]
        return x

    pad.write_text(yaml.safe_dump(schoon(record), allow_unicode=True, sort_keys=False, width=90),
                   encoding="utf-8")


# ---------------- inlezen ----------------

def lees_hoofd(pad):
    wb = openpyxl.load_workbook(pad, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    H = list(rows[0])
    idx = {h: i for i, h in enumerate(H)}

    def cel(r, h):
        i = idx[h]
        return r[i] if i < len(r) else None

    def dedup(idcol, prefix):
        uit = {}
        velden = [h for h in H if h.startswith(prefix)]
        for r in rows[1:]:
            rid = cel(r, idcol)
            if rid in (None, "") or rid in uit:
                continue
            uit[rid] = {h[len(prefix):]: cel(r, h) for h in velden}
        return uit

    reizigers = {}
    eigen = H[:30]
    for r in rows[1:]:
        rid = cel(r, "Id")
        if rid in (None, "") or rid in reizigers:
            continue
        reizigers[rid] = {h: cel(r, h) for h in eigen}
    tabellen = {
        "reizigers": reizigers,
        "reizen": dedup("Reizen::Id", "Reizen::"),
        "afbeeldingen": dedup("Afbeeldingen::Id", "Afbeeldingen::"),
        "correspondentie": dedup("Correspondentie::Id", "Correspondentie::"),
        "poezie": dedup("Poëzie::Id", "Poëzie::"),
        "reisgezellen": dedup("Reisgezellen::Id", "Reisgezellen::"),
        "educatie": dedup("Educatie::Id", "Educatie::"),
    }
    return tabellen


def lees_addenda(pad):
    wb = openpyxl.load_workbook(pad, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    H = list(rows[0])
    uit = []
    for r in rows[1:]:
        d = {H[i]: (r[i] if i < len(r) else None) for i in range(len(H))}
        if any(v not in (None, "") for v in d.values()):
            uit.append(d)
    return uit


# ---------------- omzetten ----------------

def hoofd():
    if len(sys.argv) < 3:
        sys.exit("gebruik: python3 scripts/import_filemaker.py <reizigers.xlsx> <addenda.xlsx>")
    t = lees_hoofd(sys.argv[1])
    addenda = lees_addenda(sys.argv[2])

    # id-toewijzing voor reizigers en reizen (stabiel, leesbaar)
    reiz_id = {}
    gebruikt = set()
    for fid, v in t["reizigers"].items():
        naam = tekst(v.get("c_Naam")) or f"{tekst(v.get('Voornaam'))} {tekst(v.get('Achternaam'))}"
        s = slug(naam) or f"reiziger-{fid}"
        basis = s
        n = 2
        while s in gebruikt:
            s = f"{basis}-{n}"; n += 1
        gebruikt.add(s)
        reiz_id[fid] = s

    reis_id = {}
    for fid, v in t["reizen"].items():
        rfk = v.get("ReizigerIdFk")
        stam = reiz_id.get(rfk, f"reis-{fid}")
        jaren = slug(tekst(v.get("Jaren")) or tekst(v.get("DatumAanvang")) or str(fid))
        s = f"{stam}-{jaren}"
        basis = s
        n = 2
        while s in gebruikt:
            s = f"{basis}-{n}"; n += 1
        gebruikt.add(s)
        reis_id[fid] = s

    # kinderen groeperen op ouder
    per_reiziger = defaultdict(lambda: defaultdict(list))
    per_reis = defaultdict(lambda: defaultdict(list))
    for rec in t["afbeeldingen"].values():
        per_reiziger[rec.get("ReizigersIdFk")]["afbeeldingen"].append(rec)
    for rec in t["educatie"].values():
        per_reiziger[rec.get("ReizigersIdFk")]["educatie"].append(rec)
    for rec in t["correspondentie"].values():
        per_reis[rec.get("ReisIdFk")]["correspondentie"].append(rec)
    for rec in t["poezie"].values():
        per_reis[rec.get("ReisIdFk")]["poezie"].append(rec)
    for rec in t["reisgezellen"].values():
        per_reis[rec.get("ReisIdFk")]["reisgezellen"].append(rec)
    for rec in addenda:
        per_reis[rec.get("ReisIdFk")]["addenda"].append(rec)

    # ---- reizigers ----
    for fid, v in t["reizigers"].items():
        rid = reiz_id[fid]
        achter = tekst(v.get("Achternaam"))
        tussen = tekst(v.get("Tussenvoegsel"))
        voor = tekst(v.get("Voornaam")) or tekst(v.get("Doopnamen"))
        lemma = achter + (f", {voor}" if voor else "")
        if tussen:
            lemma += f" {tussen}"
        varianten = lijst(v.get("AlternatieveNamen"))
        if tekst(v.get("Doopnamen")) and tekst(v.get("Doopnamen")) != voor:
            varianten.append("Doopnamen: " + tekst(v.get("Doopnamen")))

        typering = " ".join(x for x in [tekst(v.get("EersteReis")),
                    (f"Eerste reis: {tekst(v.get('DatumEersteReis'))}" if tekst(v.get("DatumEersteReis")) else ""),
                    (f"({tekst(v.get('LeeftijdEersteReis'))} jaar)" if tekst(v.get("LeeftijdEersteReis")) else "")] if x)

        portretten = []
        for a in per_reiziger[fid]["afbeeldingen"]:
            portretten.append({
                "titel": tekst(a.get("Titel")), "geportretteerde": tekst(a.get("Geportretteerde")),
                "vervaardiger": tekst(a.get("Vervaardiger")), "jaar": tekst(a.get("Jaar")),
                "bron": tekst(a.get("Bron")), "link": tekst(a.get("Link")),
                "beschrijving": tekst(a.get("Beschrijving"))})
        opleiding = []
        for e in per_reiziger[fid]["educatie"]:
            opleiding.append({
                "universiteit": tekst(e.get("Universiteit")), "richting": tekst(e.get("Richting")),
                "datum": tekst(e.get("Datum")), "incipit": tekst(e.get("Incipit")),
                "opmerking": tekst(e.get("Opmerking")), "bron": tekst(e.get("Bron"))})

        genealogie = {
            "vader": tekst(v.get("Vader")), "moeder": tekst(v.get("Moeder")),
            "echtgenoten": lijst(v.get("Echtgenoten")), "kinderen": lijst(v.get("Kinderen")),
            "broers_zussen": lijst(v.get("BroersZussen")), "andere_familie": lijst(v.get("AndereFamilie"))}

        schrijf("reizigers", {
            "id": rid, "naam": lemma, "volledige_naam": tekst(v.get("c_Naam")),
            "naamsvarianten": varianten, "titel": tekst(v.get("Titel")),
            "achtervoegsel": tekst(v.get("Achtervoegsel")),
            "geslacht": tekst(v.get("Geslacht")),
            "geboren": {**(datumveld(v.get("Geboortedatum")) or {}), **plaatsdict(v.get("Geboorteplaats"), v.get("Geboorteprovincie"))} or None,
            "overleden": {**(datumveld(v.get("Sterfdatum")) or {}), **plaatsdict(v.get("Sterfplaats"), v.get("Sterfprovincie"))} or None,
            "religie": tekst(v.get("Religie")),
            "is_auteur": tekst(v.get("Auteur")) == "Auteur",
            "korte_typering": typering,
            "biografie": tekst(v.get("Biografie")),
            "genealogie": genealogie,
            "opleiding": opleiding,
            "portretten": portretten,
            "literatuur": lijst(v.get("Literatuur")),
            "bronnen": [BRON],
        })

    # ---- reizen + manuscripten ----
    for fid, v in t["reizen"].items():
        rid = reis_id[fid]
        rfk = v.get("ReizigerIdFk")
        auteur = reiz_id.get(rfk)
        reistype = REISTYPE.get(tekst(v.get("Doel")), "onbekend")
        jaren = tekst(v.get("Jaren"))
        titel = tekst(v.get("Titel"))
        naam_reiziger = tekst(t["reizigers"].get(rfk, {}).get("c_Naam")) or "onbekende reiziger"
        reistitel = f"{REISTYPE_TITEL[reistype]} van {naam_reiziger}" + (f" ({jaren})" if jaren else "")

        gezellen = []
        for g in per_reis[fid]["reisgezellen"]:
            nm = " ".join(x for x in [tekst(g.get("Voornaam")), tekst(g.get("Tussenvoegsel")),
                          tekst(g.get("Achternaam")), tekst(g.get("TitelAchtervoegsel"))] if x)
            jr = "–".join(x for x in [tekst(g.get("Geboortejaar")), tekst(g.get("Sterfjaar"))] if x)
            gezellen.append({"naam": nm.strip(), "jaren": jr, "opmerkingen": tekst(g.get("Opmerkingen"))})

        brieven = []
        for b in per_reis[fid]["correspondentie"]:
            brieven.append({
                "verzender": tekst(b.get("Verzender")), "ontvanger": tekst(b.get("Ontvanger")),
                "datum": tekst(b.get("Datum")),
                "plaats_verzender": tekst(b.get("PlaatsVerzender")), "land_verzender": tekst(b.get("LandVerzender")),
                "plaats_ontvanger": tekst(b.get("PlaatsOntvanger")), "land_ontvanger": tekst(b.get("LandOntvanger")),
                "relatie": tekst(b.get("Relatie")), "taal": tekst(b.get("Taal")),
                "bron": tekst(b.get("Bron")), "link": tekst(b.get("Link")),
                "opmerking": tekst(b.get("Opmerking")), "tekst": tekst(b.get("Tekst"))})
        gedichten = []
        for p in per_reis[fid]["poezie"]:
            gedichten.append({
                "titel": tekst(p.get("Titel")), "dichter": tekst(p.get("Dichter")),
                "aan": tekst(p.get("Aan")), "datum": tekst(p.get("Datum")),
                "locatie": tekst(p.get("Locatie")), "taal": tekst(p.get("Taal")),
                "bron": tekst(p.get("Bron")), "link": tekst(p.get("Link")),
                "beschrijving": tekst(p.get("Beschrijving")), "tekst": tekst(p.get("Tekst"))})
        addenda_lijst = []
        for a in per_reis[fid]["addenda"]:
            addenda_lijst.append({
                "type": tekst(a.get("Type")), "auteur": tekst(a.get("Auteur")),
                "beschrijving": tekst(a.get("Beschrijving")), "datum": tekst(a.get("Datum")),
                "relatie": tekst(a.get("Relatie")), "taal": tekst(a.get("Taal")),
                "bron": tekst(a.get("Bron")), "link": tekst(a.get("Link")),
                "locatie": tekst(a.get("Locatie")),
                "transcriptie": tekst(a.get("Transcriptie")), "foto": tekst(a.get("Foto"))})

        heeft_ms = any(tekst(v.get(k)) for k in ("Signatuur", "Bibliotheek", "Folia", "Pagina's", "OngenummerdPagina's"))
        ms_id = f"{rid}-hs" if heeft_ms else None

        route = []
        for kant in ("Aanvang", "Einde"):
            pl = tekst(v.get(f"Plaats{kant}"))
            if pl:
                route.append({"plaats": pl, "land_modern": tekst(v.get(f"Land{kant}")),
                              "datum": tekst(v.get(f"Datum{kant}"))})

        schrijf("reizen", {
            "id": rid, "titel": titel or reistitel,
            "reizigers": [auteur] if auteur else [],
            "reistype": reistype, "jaren": jaren,
            "vertrek": tekst(v.get("DatumAanvang")) or (jaren.split("-")[0] if jaren else ""),
            "terugkeer": tekst(v.get("DatumEinde")),
            "reisdagen": tekst(v.get("Reisdagen")),
            "gebieden": lijst(v.get("Landen")),
            "route": route,
            "manuscript": ms_id,
            "beschrijving": tekst(v.get("Beschrijving")),
            "reisgezellen": gezellen,
            "brieven": brieven,
            "poezie": gedichten,
            "addenda": addenda_lijst,
            "literatuur": lijst(v.get("Literatuur")),
            "bronnen": [BRON],
        })

        if heeft_ms:
            omvang = []
            folia = tekst(v.get("Folia"))
            paginas = tekst(v.get("Pagina's"))
            if folia:
                omvang.append(f"ff. {folia}")
            if paginas:
                omvang.append(f"pp. {paginas}")
            bib = tekst(v.get("Bibliotheek"))
            iid = instelling_id(v.get("Bibliotheek"), v.get("Locatie"))
            opm = []
            if "/" in bib:
                opm.append("Meerdere bewaarplaatsen: " + bib + ".")
            schrijf("manuscripten", {
                "id": ms_id, "titel_aanduiding": titel or f"Handschrift van {reistitel}",
                "reizen": [rid], "auteurs": [auteur] if auteur else [],
                "instelling": iid, "signatuur": tekst(v.get("Signatuur")),
                "taal": tekst(v.get("Taal")).replace("\n", ", "),
                "omvang": "; ".join(omvang),
                "incompleet": tekst(v.get("Incompleet")) == "Ja",
                "beschrijving": tekst(v.get("Beschrijving")),
                "corpus_gor": tekst(v.get("Proef.No")),
                "repertorium_lsd": tekst(v.get("LSD")),
                "edities": lijst(v.get("Editie")),
                "digitalisering": {"status": "onbekend"},
                "transcriptie": {"status": "onbekend"},
                "opmerkingen": " ".join(opm),
                "bronnen": [BRON],
            })

    for rec in instellingen_gebruikt.values():
        schrijf("instellingen", rec)

    n_ms = len(list((DATA / "manuscripten").glob("*.yaml")))
    print(f"Geschreven: {len(t['reizigers'])} reizigers, {len(t['reizen'])} reizen, "
          f"{n_ms} manuscripten, {len(instellingen_gebruikt)} instellingen.")
    print(f"Verrijkingen genest: {len(t['afbeeldingen'])} portretten, "
          f"{len(t['correspondentie'])} brieven, {len(t['poezie'])} gedichten, "
          f"{len(addenda)} addenda, {len(t['reisgezellen'])} reisgezellen, {len(t['educatie'])} opleidingen.")
    if onbekende_instellingen:
        print("Nieuw aangemaakte instellingen (controle):", "; ".join(sorted(onbekende_instellingen)))


if __name__ == "__main__":
    hoofd()
